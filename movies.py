import streamlit as st
import pandas as pd
import time
from selenium import webdriver
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.chrome.options import Options
from webdriver_manager.chrome import ChromeDriverManager
from bs4 import BeautifulSoup
import re
import base64
from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

st.set_page_config(page_title="IMDb Top 250 Movies", page_icon="🎬", layout="wide")

st.markdown("""
<style>
body { background: #0e0e0e; color: #ffffff; }
.main-title {
    text-align: center; font-size: 45px; font-weight: 800;
    color: #FFD700; text-shadow: 2px 2px 10px #000; margin-top: -20px;
}
.subtitle {
    text-align: center; font-size: 20px; color: #cccccc;
    margin-top: -10px; margin-bottom: 30px;
}
.movie-card {
    background: rgba(255,255,255,0.08); border-radius: 15px;
    padding: 15px; text-align: center; margin-bottom: 15px;
    transition: all 0.3s ease;
}
.movie-card:hover { transform: scale(1.05); box-shadow: 0 0 25px rgba(255, 215, 0, 0.3); }
a { text-decoration: none; }
::-webkit-scrollbar { width: 8px; }
::-webkit-scrollbar-thumb { background: #FFD700; border-radius: 10px; }
</style>
""", unsafe_allow_html=True)

st.markdown("<div class='main-title'>🎬 IMDb Top 250 Movies Explorer</div>", unsafe_allow_html=True)
st.markdown("<div class='subtitle'>View, analyze, and download the most popular films in IMDb history</div>", unsafe_allow_html=True)

@st.cache_data(show_spinner=False)
def scrape_imdb_top250():
    options = Options()
    options.add_argument("--start-maximized")
    options.add_argument("--disable-blink-features=AutomationControlled")
    options.add_experimental_option("detach", True)
    options.add_argument("--no-sandbox")
    options.add_argument("--disable-dev-shm-usage")

    driver = webdriver.Chrome(service=Service(ChromeDriverManager().install()), options=options)
    driver.get("https://www.imdb.com/chart/top/")

    last_height = driver.execute_script("return document.body.scrollHeight")
    while True:
        driver.execute_script("window.scrollTo(0, document.body.scrollHeight);")
        time.sleep(2)
        new_height = driver.execute_script("return document.body.scrollHeight")
        if new_height == last_height:
            break
        last_height = new_height

    html = driver.page_source
    soup = BeautifulSoup(html, "html.parser")

    rows = soup.select("li.ipc-metadata-list-summary-item")
    movies = []
    for row in rows:
        try:
            title_raw = row.select_one("h3").get_text(strip=True)
            title = re.sub(r"^\d+\.\s*", "", title_raw)
            rating_tag = row.select_one("span.ipc-rating-star--rating")
            rating = rating_tag.get_text(strip=True) if rating_tag else "N/A"
            year_tag = row.select_one("span.ipc-metadata-list-summary-item__li")
            year = year_tag.get_text(strip=True) if year_tag else "N/A"
            year = re.findall(r"\d{4}", year)[0] if re.findall(r"\d{4}", year) else ""
            link = "https://www.imdb.com" + row.select_one("a")["href"]
            img_tag = row.select_one("img")
            poster = img_tag["src"] if img_tag else ""
            movies.append([title, rating, year, link, poster])
        except:
            continue

    df = pd.DataFrame(movies, columns=["Title", "Rating", "Year", "Link", "Poster"])
    df["Rating"] = pd.to_numeric(df["Rating"], errors="coerce")
    return df

with st.spinner("🚀 Extracting IMDb Top 250 movies dynamically..."):
    df = scrape_imdb_top250()

left_col, right_col = st.columns([1, 2])

with left_col:
    st.markdown("### 💾 Export Dataset")

    excel_file = "IMDb_Top_250.xlsx"
    df.to_excel(excel_file, index=False)

    wb = load_workbook(excel_file)
    ws = wb.active

    for row in range(2, len(df) + 2):
        title_cell = ws[f"A{row}"]
        url = ws[f"D{row}"].value
        title_cell.value = f'=HYPERLINK("{url}", "{df.iloc[row-2, 0]}")'
        title_cell.font = Font(color="0000FF", underline="single")

    for col in ws.columns:
        max_length = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = max_length + 8
        if col_letter == "A":
            adjusted_width = 60
        ws.column_dimensions[col_letter].width = adjusted_width

    wb.save(excel_file)

    with open(excel_file, "rb") as f:
        b64 = base64.b64encode(f.read()).decode()

    st.markdown(f"""
        <style>
        .download-button {{
            display: inline-flex;
            align-items: center;
            justify-content: center;
            background-color: #4DA8FF;
            color: white;
            font-weight: 700;
            border-radius: 12px;
            padding: 12px 22px;
            font-size: 16px;
            text-decoration: none;
            transition: 0.3s ease-in-out;
        }}
        .download-button:hover {{
            background-color: #1E90FF;
            transform: scale(1.05);
        }}
        .download-icon {{
            margin-right: 8px;
            font-size: 18px;
        }}
        </style>

        <div style="text-align:center;">
            <a href="data:application/octet-stream;base64,{b64}" download="{excel_file}" class="download-button">
                <span class="download-icon">📥</span>Download as Excel
            </a>
        </div>
    """, unsafe_allow_html=True)

    st.markdown("### 🔍 Search / Filter Movies")

    search_term = st.text_input("Search by Title").lower()
    min_rating = st.slider("Minimum IMDb Rating", 0.0, 10.0, 0.0, 0.1)
    year_filter = st.text_input("Release Year (e.g. 2024)")

    filtered_df = df.copy()
    if search_term:
        filtered_df = filtered_df[filtered_df["Title"].str.lower().str.contains(search_term, na=False)]
    if min_rating > 0:
        filtered_df = filtered_df[filtered_df["Rating"] >= min_rating]
    if year_filter.strip():
        filtered_df = filtered_df[filtered_df["Year"].astype(str).str.fullmatch(year_filter.strip())]

    st.markdown(f"📊 Results: {len(filtered_df)} of {len(df)} Movies**")

with right_col:
    num_cols = 4
    cols = st.columns(num_cols)
    for idx, row in filtered_df.iterrows():
        with cols[idx % num_cols]:
            st.markdown(f"""
                <div class='movie-card'>
                    <img src='{row["Poster"]}' width='150'><br>
                    <b>{row["Title"]}</b><br>
                    ⭐ {row["Rating"] if not pd.isna(row["Rating"]) else "N/A"} | 📅 {row["Year"]}<br>
                    <a href='{row["Link"]}' target='_blank' style='color:#4DA8FF;'>🎬 View IMDb Page</a>
                </div>
            """, unsafe_allow_html=True)
        if (idx + 1) % num_cols == 0:
            cols = st.columns(num_cols)